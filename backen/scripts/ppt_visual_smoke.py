#!/usr/bin/env python3
"""Render every built-in PPT style through python-pptx and LibreOffice.

This is intentionally deterministic: it exercises the renderer's real style
branches with a small five-slide deck and fails if LibreOffice cannot open the
result or the slide count changes. It is a release smoke check, not a content
quality benchmark.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation


STYLES = [
    ("academic-blue", "academic", ["005BAC", "063A78", "D9A441", "EFF6FF", "1F2937"]),
    ("minimal-ink", "minimal", ["111827", "374151", "0EA5E9", "F8FAFC", "1F2937"]),
    ("emerald-report", "report", ["047857", "064E3B", "F59E0B", "ECFDF5", "1F2937"]),
    ("warm-defense", "warm", ["B45309", "7C2D12", "2563EB", "FFF7ED", "1F2937"]),
    ("gaia-editorial", "editorial", ["9A3412", "431407", "0F766E", "FFF7ED", "292524"]),
    ("uncover-contrast", "contrast", ["0F172A", "020617", "F97316", "F8FAFC", "E2E8F0"]),
    ("dracula-night", "dark-tech", ["BD93F9", "282A36", "50FA7B", "282A36", "F8F8F2"]),
    ("slidev-seriph", "seriph", ["2563EB", "172554", "F59E0B", "F8FAFC", "334155"]),
    ("slidev-apple-basic", "apple-basic", ["111827", "000000", "3B82F6", "FFFFFF", "374151"]),
    ("startup-pitch", "pitch", ["7C3AED", "312E81", "F59E0B", "F5F3FF", "1F2937"]),
    ("product-launch", "product", ["0E7490", "164E63", "F43F5E", "ECFEFF", "164E63"]),
    ("training-canvas", "training", ["2563EB", "1E3A8A", "F97316", "EFF6FF", "1E293B"]),
    ("ppt-master-editorial", "editorial", ["C2410C", "431407", "F59E0B", "FFF7ED", "292524"]),
    ("ppt-master-memphis", "memphis", ["F43F5E", "312E81", "FACC15", "FFF1F2", "1E1B4B"]),
    ("ppt-master-data-journalism", "data-journalism", ["38BDF8", "0F172A", "FBBF24", "111827", "E2E8F0"]),
    ("ppt-master-swiss-grid", "swiss-grid", ["DC2626", "111827", "FDE047", "F8FAFC", "1F2937"]),
    ("ppt-master-glassmorphism", "glass-saas", ["A78BFA", "111827", "22D3EE", "111827", "F8FAFC"]),
    ("presenton-glass-saas", "glass-saas", ["8B5CF6", "1E1B4B", "2DD4BF", "F5F3FF", "EDE9FE"]),
    ("presenton-gradient-pitch", "gradient-pitch", ["F97316", "4C1D95", "FDE68A", "F5F3FF", "312E81"]),
    ("presenton-product-studio", "product-studio", ["06B6D4", "164E63", "FB7185", "ECFEFF", "164E63"]),
    ("primer-github-blueprint", "primer", ["0969DA", "1F2328", "54AEFF", "F6F8FA", "1F2328"]),
    ("primer-data-report", "primer-report", ["8250DF", "24292F", "BF8700", "FFFFFF", "24292F"]),
    ("primer-open-source", "primer-open-source", ["1A7F37", "24292F", "9A6700", "F6F8FA", "24292F"]),
]


def deck_for(design: str) -> dict:
    return {
        "title": f"{design} smoke test",
        "subtitle": "Deterministic renderer visual contract",
        "theme": design,
        "templateDesign": design,
        "slides": [
            {"type": "cover", "title": "封面标题", "headline": "五页视觉回归测试"},
            {"type": "section", "section": "CONTEXT", "title": "章节页", "headline": "验证章节节奏"},
            {"type": "content", "section": "EVIDENCE", "title": "图文内容页", "headline": "验证正文、指标与留白", "bullets": ["短要点一", "短要点二", "短要点三"], "metrics": [{"value": "98%", "label": "覆盖率"}]},
            {"type": "content", "section": "COMPARISON", "title": "对比页", "layout": "comparison", "bullets": ["基线方案", "改进方案", "证据与结论"], "metrics": [{"value": "2x", "label": "速度"}, {"value": "-30%", "label": "成本"}]},
            {"type": "conclusion", "section": "OUTLOOK", "title": "结论与下一步", "headline": "所有内置风格都应可打开、可渲染、可继续编辑", "bullets": ["输出可打开", "版式不溢出", "颜色继承模板"]},
        ],
    }


def run(command: list[str], cwd: Path | None = None) -> None:
    result = subprocess.run(command, cwd=cwd, text=True, capture_output=True)
    if result.returncode != 0:
        raise RuntimeError(f"命令失败: {' '.join(command)}\n{result.stdout}\n{result.stderr}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--renderer", default=str(Path(__file__).with_name("ppt_renderer.py")))
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--skip-libreoffice", action="store_true")
    args = parser.parse_args()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not args.skip_libreoffice and not soffice:
        raise SystemExit("找不到 LibreOffice/soffice；如只做结构 smoke，请显式传 --skip-libreoffice")

    root = (args.output_dir or Path(tempfile.mkdtemp(prefix="ppt-visual-smoke-"))).resolve()
    root.mkdir(parents=True, exist_ok=True)
    results = []
    try:
        for key, design, palette in STYLES:
            work = root / key
            work.mkdir(parents=True, exist_ok=True)
            deck_path = work / "deck.json"
            style_path = work / "style.json"
            images_path = work / "images.json"
            manifest_path = work / "manifest.json"
            output = work / "output.pptx"
            deck_path.write_text(json.dumps(deck_for(design), ensure_ascii=False), encoding="utf-8")
            style_path.write_text(json.dumps({"palette": palette}, ensure_ascii=False), encoding="utf-8")
            images_path.write_text("[]", encoding="utf-8")
            manifest_path.write_text('{"images": []}', encoding="utf-8")
            run([sys.executable, args.renderer, "--deck", str(deck_path), "--style", str(style_path),
                 "--images", str(images_path), "--manifest", str(manifest_path), "--out", str(output)])
            prs = Presentation(str(output))
            if len(prs.slides) != 5:
                raise RuntimeError(f"{key}: 期望 5 页，实际 {len(prs.slides)} 页")
            pdf = work / "output.pdf"
            if not args.skip_libreoffice:
                profile = work / "lo-profile"
                run([soffice, "--headless", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf",
                     "--outdir", str(work), str(output)])
                if not pdf.is_file() or pdf.stat().st_size < 1024:
                    raise RuntimeError(f"{key}: LibreOffice 未生成有效 PDF")
            results.append({"key": key, "slides": len(prs.slides), "pdf": pdf.is_file()})
        print(json.dumps({"ok": True, "styles": results, "outputDir": str(root)}, ensure_ascii=False, indent=2))
        return 0
    finally:
        if args.output_dir is None:
            shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
